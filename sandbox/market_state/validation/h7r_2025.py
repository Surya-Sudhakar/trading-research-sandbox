"""Execution-controlled H7 replication on exposed 2025 data; explicit invocation only.

The sole parser adaptation is EURUSD_DUKASCOPY_TICK -> EURUSD_DUKASCOPY in
validated symbol cells of an in-memory copy. Original H7 code and inputs are immutable.
"""
from copy import deepcopy
from dataclasses import asdict
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import re

from openpyxl import load_workbook

from sandbox.market_state.validation import h7_holdout_2025 as h7
from sandbox.provenance import sha256_file

EXECUTION_SYMBOL = 'EURUSD_DUKASCOPY_TICK'
PARSER_SYMBOL = 'EURUSD_DUKASCOPY'
REPORT = Path('data/external/dukascopy/EURUSD/h7r_2025_genuine_ticks.xlsx')
REPORT_SHA256 = 'e68201e219d8bd51182dc3974d0f2cf175bb576ebe790d3e726c9afb1238f379'
CSV_SHA256 = '9dc62c75808e24b84baf2965ee6a34c56a192e3da4ca882850561b4b624f960e'
H7_CODE_SHA256 = '5bf972047acf2b80e95143b6bed0f73483cd031221d6b32ac56858ffc9985f91'
OUTPUT = Path('results/market_state')
NAMES = ('h7r_2025_trades.csv','h7r_2025_summary.json','h7r_2025.attempt.json')


def _value(rows, label):
    found = []
    for row in rows:
        for i, cell in enumerate(row):
            if cell == label:
                found.append(next((v for v in row[i+1:] if v is not None), None))
    if len(found) != 1 or found[0] is None:
        raise ValueError('missing/ambiguous report field: '+label)
    return found[0]


def adapt_symbols(rows):
    """Strict, H7-R-specific accommodation; never edit workbook or H7 globals."""
    original = [tuple(row) for row in rows]
    if _value(original,'Symbol:') != EXECUTION_SYMBOL:
        raise ValueError('H7-R requires exactly '+EXECUTION_SYMBOL)
    order_headers = [i for i,r in enumerate(original) if r[:4] == ('Open Time','Order','Symbol','Type')]
    deal_headers = [i for i,r in enumerate(original) if r[:5] == ('Time','Deal','Symbol','Type','Direction')]
    if len(order_headers)!=1 or len(deal_headers)!=1 or order_headers[0]>=deal_headers[0]:
        raise ValueError('unique Orders and Deals tables required')
    result = [list(row) for row in original]
    for i,row in enumerate(original):
        if row and row[0]=='Symbol:':
            if len(row)<4 or row[3]!=EXECUTION_SYMBOL:
                raise ValueError('invalid Symbol setting layout')
            result[i][3]=PARSER_SYMBOL
        order = order_headers[0]<i<deal_headers[0] and len(row)>11 and isinstance(row[0],str) and re.fullmatch(r'\d{4}\.\d{2}\.\d{2} .*',row[0])
        deal = i>deal_headers[0] and len(row)>4 and row[3] in ('buy','sell')
        if order or deal:
            if row[2]!=EXECUTION_SYMBOL:
                raise ValueError('H7-R order/deal execution symbol mismatch')
            result[i][2]=PARSER_SYMBOL
    return result


def parse_report_rows(rows):
    rows = [tuple(row) for row in rows]
    trades, report = h7.parse_report_rows(adapt_symbols(rows))
    # The supplied execution report covers 2025. Do not compare selected-period
    # outcomes to whole-report outcomes if another period is ever supplied.
    if report['excluded_outside_2025']:
        raise ValueError('H7-R report reconciliation requires 2025-only completed trades')
    deal_header = next(i for i,r in enumerate(rows) if r[:5]==('Time','Deal','Symbol','Type','Direction'))
    entries=sum(r[3:5]==('buy','in') for r in rows[deal_header+1:])
    exits=sum(r[3:5]==('sell','out') for r in rows[deal_header+1:])
    wins=int(trades.outcome.eq('WIN').sum())
    losses=int(trades.outcome.eq('LOSS').sum())
    def count(label):
        match=re.fullmatch(r'(\d+) \([\d.]+%\)',str(_value(rows,label)))
        if not match:
            raise ValueError('invalid report outcome count: '+label)
        return int(match.group(1))
    if entries!=len(trades) or exits!=len(trades) or wins+losses!=len(trades):
        raise ValueError('entry/exit/trade reconciliation failed')
    if wins!=count('Profit Trades (% of total):') or losses!=count('Loss Trades (% of total):'):
        raise ValueError('report WIN/LOSS totals do not reconcile')
    net=sum((Decimal(str(v)) for v in trades.profit),Decimal(0))
    if net!=h7.cash(_value(rows,'Total Net Profit:')):
        raise ValueError('net trade cash does not reconcile')
    expected=net/len(trades)
    displayed=h7.cash(_value(rows,'Expected Payoff:'))
    # MT5 displays expected payoff to six decimal places. This is display
    # precision reconciliation, not a statistical or trading threshold.
    if abs(expected-displayed)>Decimal('0.0000005'):
        raise ValueError('expected payoff does not reconcile at report precision')
    report.update(execution_symbol=EXECUTION_SYMBOL,
        symbol_adapter={'required_source':EXECUTION_SYMBOL,'parser_only_alias':PARSER_SYMBOL,
                        'scope':'in-memory Symbol setting and Orders/Deals symbol cells only'},
        ea=str(_value(rows,'Expert:')),ea_inputs=list(h7.INPUTS),
        completed_trades=len(trades),entries=entries,exits=exits,wins=wins,losses=losses,
        net_profit_eur=float(net),expected_payoff_eur=float(expected),
        report_expected_payoff_eur=float(displayed),fully_reconciled=True)
    return trades,report


def read_report(path=REPORT):
    workbook=load_workbook(path,read_only=True,data_only=True)
    try:
        if len(workbook.worksheets)!=1:
            raise ValueError('single report sheet required')
        return parse_report_rows(workbook.worksheets[0].iter_rows(values_only=True))
    finally:
        workbook.close()


def verify_frozen(report=REPORT):
    expected={Path(report):REPORT_SHA256,h7.DATA:CSV_SHA256,
              h7.h5.DEFAULT_MODEL:h7.EXPECTED_MODEL_SHA256,
              Path(h7.__file__):H7_CODE_SHA256,h7.h5.EA:h7.h5.EA_SHA}
    for path,digest in expected.items():
        if sha256_file(path)!=digest:
            raise ValueError('frozen input SHA256 mismatch: '+str(path))
    return {str(p):digest for p,digest in expected.items()}


def protected_snapshot():
    paths={p for p in OUTPUT.rglob('*') if p.is_file()}
    paths.update(Path('sandbox/market_state').rglob('*.py'))
    paths.update(Path('sandbox/strategies/version0').rglob('*.py'))
    paths.update([h7.REPORT,h7.DATA,h7.h5.EA,Path('sandbox/partition/service.py'),
                  Path('sandbox/market_data/gaps.py'),Path('sandbox/market_data/sanitation.py')])
    return {str(p):sha256_file(p) for p in sorted(paths)}


def assert_unchanged(snapshot):
    for path,digest in snapshot.items():
        if sha256_file(Path(path))!=digest:
            raise RuntimeError('protected input/artifact changed: '+path)


def _write_new(path, payload):
    with path.open('x',encoding='utf-8',newline='') as stream:
        stream.write(payload)


def run(report=REPORT, output=OUTPUT):
    """Explicit future execution; importing or reading the report never replays states."""
    output=Path(output)
    targets=[output/name for name in NAMES]
    if any(p.exists() for p in targets):
        raise FileExistsError('H7-R attempt/output already exists; refusing overwrite')
    before=protected_snapshot()
    before.update(verify_frozen(report))
    with h7.no_fitting():
        contract,model,contract_sha=h7.load_frozen_inputs()
        profile,policy,continuity=h7.frozen_continuity(contract)
    feature_path=Path(contract['feature_artifact']['path'])
    meta_path=feature_path.with_name('metadata.json')
    for path,digest in ((feature_path,contract['feature_artifact']['sha256']),
                         (meta_path,contract['feature_artifact']['metadata_sha256'])):
        if sha256_file(path)!=digest:
            raise ValueError('frozen feature provenance mismatch')
        before[str(path)]=digest
    meta=json.loads(meta_path.read_text())
    if meta['definitions']!=[asdict(d) for d in h7.definitions('M15')]:
        raise ValueError('frozen feature definitions mismatch')
    protocol=deepcopy(h7.PROTOCOL)
    protocol.update(version='H7R_2025_V1',data_status='EXPOSED_2025',execution_symbol=EXECUTION_SYMBOL,
                    replication_of='H7_HOLDOUT_2025_V1',
                    experimental_change='MT5 execution report from genuine Bid/Ask tick symbol; state input unchanged')
    attempt={'status':'RUNNING','protocol':protocol,'execution_symbol':EXECUTION_SYMBOL,
             'report_path':str(report),'report_sha256':REPORT_SHA256,
             'protected_sha256_before':before,'implementation_sha256':sha256_file(Path(__file__))}
    _write_new(targets[2],json.dumps(attempt,indent=2,sort_keys=True)+'\n')
    try:
        with h7.no_fitting() as guard:
            trades,report_summary=read_report(report)
            bars,_,data=h7.read_bars(h7.DATA,profile,policy)
            if trades.entry_time.min()<bars.timestamp_utc.min() or trades.exit_time.max()>bars.timestamp_utc.max()+h7.STEP:
                raise ValueError('deal timestamps outside source coverage')
            features,states=h7.state_rows(bars,model)
            joined=h7.attach(trades,states,features,bars,meta['definitions'])
            summary=h7.summarize(joined)
        if guard['fit_calls']:
            raise RuntimeError('fit attempted')
        assert_unchanged(before)
        summary.update(experiment=protocol['version'],protocol=protocol,data=data,MT5_report=report_summary,
            continuity=continuity,feature_rows=len(features),complete_state_rows=len(states),
            provenance={'execution_symbol':EXECUTION_SYMBOL,'report_path':str(report),
                'report_sha256':REPORT_SHA256,'csv_sha256':CSV_SHA256,
                'model_sha256':h7.EXPECTED_MODEL_SHA256,'contract_file_sha256':contract_sha,
                'contract_canonical_sha256':h7.CONTRACT_SHA256,'ea_sha256':h7.h5.EA_SHA,
                'protected_sha256_before_and_after':before,'fit_calls':guard['fit_calls']},
            limitations=['2025 is exposed data; this is execution-controlled replication, not a new holdout.',
                'Same frozen H7 comparisons and pointwise day-cluster bootstrap; no new materiality cutoff.',
                'Incomplete vectors remain unmatched; no added warm-up, imputation or stale-state carry.',
                'History-quality text is report metadata, not independent tick-source certification.'])
        csv=joined.to_csv(index=False,lineterminator='\n')
        summary['provenance']['trades_csv_sha256']=hashlib.sha256(csv.encode()).hexdigest()
        payload=json.dumps(summary,indent=2,sort_keys=True,allow_nan=False)+'\n'
        _write_new(targets[0],csv)
        _write_new(targets[1],payload)
        assert_unchanged(before)
        attempt.update(status='COMPLETED',fit_calls=guard['fit_calls'],
                       outputs={str(p):sha256_file(p) for p in targets[:2]})
    except Exception as error:
        attempt.update(status='FAILED',error_type=type(error).__name__,error=str(error))
        raise
    finally:
        targets[2].write_text(json.dumps(attempt,indent=2,sort_keys=True)+'\n',encoding='utf-8')
    return summary


if __name__=='__main__':
    result=run()
    print(json.dumps({k:result[k] for k in ('experiment','overall','reconciliation','by_state')},indent=2))
