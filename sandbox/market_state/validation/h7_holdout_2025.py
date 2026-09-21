"""One-shot 2025 H7: actual MT5 deals, frozen causal features and frozen inference."""
from collections import Counter
from dataclasses import asdict
from datetime import timedelta
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import re
from types import SimpleNamespace

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from sandbox.market_data.sanitation import sanitize
from sandbox.market_state.universal.schema import definitions
from sandbox.market_state.discovery import h5_discovery_2016_2022 as h5
from sandbox.market_state.discovery.cross_feed_diagnostic import infer, FEATURES, POSTERIORS
from sandbox.market_state.discovery.dukascopy_h5_coverage import frozen_continuity
from sandbox.market_state.discovery.future_behavior import _effect
from sandbox.market_state.discovery.validation import load_frozen_inputs, no_fitting, EXPECTED_MODEL_SHA256
from sandbox.market_state.discovery.validation_contract import CONTRACT_SHA256
from sandbox.provenance import sha256_file

START = pd.Timestamp('2025-01-01', tz='UTC')
END = pd.Timestamp('2026-01-01', tz='UTC')
STEP = timedelta(minutes=15)
STATES = ('S0', 'S1', 'S2')
DATA = Path('data/external/dukascopy/EURUSD/eurusd-m15-bid-2025-01-01-2026-01-01.csv')
REPORT = Path('data/external/dukascopy/EURUSD/2025-2026-dukas.xlsx')
OUTPUT = Path('results/market_state')
INPUTS = ('TakeProfitPips=5', 'StopLossPips=40', 'PipSize=0.0001', 'FixedLots=0.01',
          'MagicNumber=26082930', 'EnableTrading=true', 'UseManualServerUTCOffset=true',
          'ManualServerUTCOffsetHours=0')
PROTOCOL = {
    'version': 'H7_HOLDOUT_2025_V1', 'start_inclusive': START.isoformat(), 'end_exclusive': END.isoformat(),
    'population': 'actual completed MT5 Version 0 trades whose entries and exits are in 2025; no synthetic trades',
    'hypothesis': 'S1 has better conditional Version 0 expectancy than S0 and S2',
    'comparisons': [['S1', 'S0'], ['S1', 'S2']],
    'directional_check': 'S1 mean net EUR profit/trade > S0 AND > S2; independent of overall profitability',
    'materiality': 'No numerical materiality threshold was specified; report effect sizes and uncertainty without inventing one',
    'statistics': 'net-EUR expectancy differences, win-rate differences, pooled-SD Cohen d',
    'uncertainty': {'method': 'whole UTC entry-day cluster bootstrap, shared resamples across states',
                    'seed': 42, 'replicates': 2000, 'interval': 'pointwise percentile 95%; no tuning'},
    'eligibility': 'same 14 frozen features, no imputation, exact floor(entry UTC,15min); no stale-state carry',
    'monthly': 'all twelve UTC entry months and all three states; descriptive only',
    'model_sha256': EXPECTED_MODEL_SHA256, 'contract_canonical_sha256': CONTRACT_SHA256,
    'ea_sha256': h5.EA_SHA, 'feature_names': list(FEATURES), 'ea_inputs': list(INPUTS)}


def cash(value):
    result = Decimal(str(value))
    if not result.is_finite():
        raise ValueError('nonfinite report money')
    return result


def parse_report_rows(rows):
    """H5's sequential full-position deal pairing, with report-driven 2025 reconciliation."""
    rows = [tuple(r) for r in rows]
    def setting(name):
        found = [r[3] for r in rows if len(r) > 3 and r[0] == name]
        if len(found) != 1:
            raise ValueError(f'missing/ambiguous setting: {name}')
        return found[0]
    for name, value in {'Expert:': 'F1_NY3H_Green_MT5_Version0', 'Symbol:': 'EURUSD_DUKASCOPY', 'Currency:': 'EUR'}.items():
        if setting(name) != value:
            raise ValueError(f'noncanonical {name}')
    requested = setting('Period:')
    if not re.fullmatch(r'M15 \(2025\.01\.01 - \d{4}\.\d{2}\.\d{2}\)', requested):
        raise ValueError('unexpected requested tester period')
    found = [r[3] for r in rows if len(r) > 3 and isinstance(r[3], str) and '=' in r[3]]
    if set(found) != set(INPUTS) or len(found) != len(INPUTS):
        raise ValueError('canonical EA inputs mismatch')
    headers = [i for i,r in enumerate(rows) if r[:5] == ('Time','Deal','Symbol','Type','Direction')]
    order_headers = [i for i,r in enumerate(rows) if r[:4] == ('Open Time','Order','Symbol','Type')]
    if len(headers) != 1 or len(order_headers) != 1 or order_headers[0] >= headers[0]:
        raise ValueError('unique Orders and Deals tables required')
    orders = {}
    for r in rows[order_headers[0]+1:headers[0]]:
        if len(r) > 11 and isinstance(r[0], str) and re.fullmatch(r'\d{4}\.\d{2}\.\d{2} .*', r[0]):
            if r[1] in orders:
                raise ValueError('duplicate order ID')
            orders[r[1]] = r
    initial = balance = cash(setting('Initial Deposit:'))
    trades, seen, active, trading_times = [], set(), None, []
    previous = None
    for r in rows[headers[0]+1:]:
        if not any(x is not None for x in r):
            continue
        if len(r) < 12:
            raise ValueError('incomplete deal row')
        timestamp = pd.to_datetime(r[0], format='%Y.%m.%d %H:%M:%S', utc=True)
        if pd.isna(timestamp) or (previous is not None and timestamp < previous) or r[1] in seen:
            raise ValueError('invalid/duplicate/nonchronological deal')
        previous = timestamp
        seen.add(r[1])
        if r[3] == 'balance':
            if len(seen) != 1 or cash(r[10]) != initial or cash(r[11]) != initial or cash(r[8])+cash(r[9]) != 0:
                raise ValueError('unexpected balance operation')
            continue
        trading_times.append(timestamp)
        if r[2] != 'EURUSD_DUKASCOPY' or cash(r[5]) != Decimal('.01'):
            raise ValueError('symbol/volume mismatch')
        price = float(r[6])
        if not np.isfinite(price) or price <= 0:
            raise ValueError('invalid fill price')
        money = sum((cash(r[k]) for k in (8,9,10)), Decimal(0))
        balance += money
        if balance != cash(r[11]):
            raise ValueError('deal balance reconciliation failed')
        if (r[3],r[4]) == ('buy','in') and active is None:
            order = orders.get(r[7])
            if order is None or order[2] != 'EURUSD_DUKASCOPY' or order[3] not in ('buy','buy limit') or order[11] != 'filled':
                raise ValueError('entry deal lacks canonical filled order')
            if order[4] != '0.01 / 0.01' or cash(order[8])-cash(order[7]) != Decimal('.0045'):
                raise ValueError('filled order volume or SL/TP mismatch')
            if order[3] == 'buy limit' and (cash(order[6])-cash(order[7]) != Decimal('.004') or cash(order[8])-cash(order[6]) != Decimal('.0005')):
                raise ValueError('limit SL/TP mismatch')
            if pd.to_datetime(order[9],format='%Y.%m.%d %H:%M:%S',utc=True) != timestamp:
                raise ValueError('deal/order fill timestamp mismatch')
            active = {'entry_time':timestamp, 'entry_price':price, 'entry_deal':r[1], 'entry_order':r[7],
                      'entry_cash':money, 'entry_commission':cash(r[8]), 'entry_swap':cash(r[9])}
        elif (r[3],r[4]) == ('sell','out') and active is not None:
            profit = active.pop('entry_cash')+money
            if not profit:
                raise ValueError('zero-net trade requires explicit unchanged outcome convention')
            commission = active.pop('entry_commission')+cash(r[8])
            swap = active.pop('entry_swap')+cash(r[9])
            trades.append({**active, 'exit_time':timestamp, 'exit_price':price, 'exit_deal':r[1],
                           'profit':float(profit), 'commission':float(commission), 'swap':float(swap),
                           'outcome':'WIN' if profit > 0 else 'LOSS'})
            active = None
        else:
            raise ValueError('partial/overlapping/noncanonical deal sequence')
    if active is not None or not trades:
        raise ValueError('incomplete position ledger')
    frame = pd.DataFrame(trades)
    if (len(frame) != setting('Total Trades:') or len(trading_times) != setting('Total Deals:') or
            balance-initial != cash(setting('Total Net Profit:'))):
        raise ValueError('report total reconciliation failed')
    included = frame.entry_time.between(START,END,inclusive='left')
    if (included & ~frame.exit_time.between(START,END,inclusive='left')).any():
        raise ValueError('2025-entry position crosses frozen H7 boundary')
    selected = frame.loc[included].reset_index(drop=True)
    if selected.empty:
        raise ValueError('no completed 2025 trades')
    return selected, {'requested_period':requested, 'report_total_trades':len(frame),
        'excluded_outside_2025':int((~included).sum()), 'report_net_eur':float(balance-initial),
        'actual_first_executed_deal':min(trading_times).isoformat(), 'actual_last_executed_deal':max(trading_times).isoformat(),
        'first_2025_entry':selected.entry_time.min().isoformat(), 'last_2025_entry':selected.entry_time.max().isoformat(),
        'last_2025_exit':selected.exit_time.max().isoformat(), 'history_quality_label':setting('History Quality:'),
        'tester_bars':setting('Bars:'), 'reconciled':True,
        'ambiguous_intrabar_outcomes':None,
        'intrabar_note':'Actual MT5 deals are used unchanged; report does not certify tick-level counterfactual ambiguity.'}


def read_report(path):
    workbook = load_workbook(path,read_only=True,data_only=True)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError('single report sheet required')
        return parse_report_rows(workbook.worksheets[0].iter_rows(values_only=True))
    finally:
        workbook.close()


def read_bars(path, profile, policy):
    raw = pd.read_csv(path,dtype=str,keep_default_na=False)
    if list(raw.columns) != ['timestamp','open','high','low','close'] or raw.empty:
        raise ValueError('Dukascopy timestamp/open/high/low/close schema required')
    if not raw.timestamp.str.fullmatch(r'[0-9]+').all():
        raise ValueError('integer epoch-millisecond timestamps required')
    times = pd.to_datetime(raw.timestamp.astype('int64'),unit='ms',utc=True)
    if times.isna().any() or not times.between(START,END,inclusive='left').all():
        raise ValueError('source observations outside frozen 2025 window')
    if not times.is_monotonic_increasing or times.duplicated().any():
        raise ValueError('source timestamps not strictly increasing')
    source = raw.drop(columns='timestamp').copy()
    source.insert(0,'timestamp_utc',times)
    metadata = SimpleNamespace(provider='DUKASCOPY',symbol='EURUSD',provider_symbol='EURUSD_DUKASCOPY',
        timeframe='M15',price_type='BID',source_dataset_id='H7-'+sha256_file(path),timezone='UTC')
    bars,audit,gaps,manifest = sanitize(source,metadata,profile,policy)
    # Do not repair a new source problem or relax frozen eligibility after opening holdout.
    if manifest['excluded_row_count']:
        raise ValueError('invalid source observations: H7 requires explicit provenance review')
    for _, group in bars.groupby('continuity_segment_id',sort=False):
        if not group.timestamp_utc.diff().iloc[1:].eq(STEP).all():
            raise ValueError('frozen continuity policy leaves an internal gap; cannot silently bridge or change policy')
    if bars.empty:
        raise ValueError('no retained source bars')
    return bars,gaps,{'path':str(path),'sha256':sha256_file(path),'M15_rows':len(raw),
        'first_timestamp':times.iloc[0].isoformat(),'last_timestamp':times.iloc[-1].isoformat(),
        'sanitation':manifest, 'gap_events':gaps.to_dict('records')}


def state_rows(bars, model):
    """Existing canonical engine/inference, and H5 age bookkeeping on complete vectors."""
    features = infer(bars,model)
    # Last bar may close exactly on the exclusive endpoint; it cannot label a 2025 entry.
    features = features.loc[features.timestamp < END].copy()
    ready = features.loc[features.state.notna()].copy().reset_index(drop=True)
    states = ready[['timestamp','continuity_segment_id','state',*POSTERIORS]].rename(
        columns={'timestamp':'state_timestamp','state':'state_at_entry'})
    connected = (states.state_timestamp.diff().eq(STEP) &
                 states.continuity_segment_id.eq(states.continuity_segment_id.shift())).to_numpy()
    age = np.ones(len(states),dtype=int)
    since = np.full(len(states),np.nan)
    labels = states.state_at_entry.to_numpy()
    for i in range(1,len(states)):
        if connected[i]:
            if labels[i] != labels[i-1]:
                since[i] = 0
            else:
                age[i] = age[i-1]+1
                since[i] = since[i-1]+1
    states['state_age'], states['bars_since_transition'] = age,since
    return features,states


def attach(trades, states, features, bars, feature_definitions):
    if not trades.entry_time.between(START,END,inclusive='left').all():
        raise ValueError('H7 trades outside 2025')
    if not states.state_timestamp.between(START,END,inclusive='left').all():
        raise ValueError('H7 states outside 2025')
    joined = h5.attach(trades,states)
    lookup = features.set_index('timestamp')
    segment_starts = bars.groupby('continuity_segment_id').timestamp_utc.min()
    requirements = {d['name']:d for d in feature_definitions}
    reasons, missing = [], []
    for row in joined.itertuples():
        if row.state_match_status == 'MATCHED':
            reasons.append('MATCHED');missing.append('[]');continue
        t = row.required_state_timestamp
        if t not in lookup.index:
            reasons.append('NO_FEATURE_ROW_AT_BOUNDARY');missing.append(None);continue
        saved = lookup.loc[t]
        names = [name for name in FEATURES if pd.isna(saved[name])]
        opened = segment_starts.loc[saved.continuity_segment_id]
        available = {'M15':int((t-opened)//STEP)}
        for tf,hours in (('H1',1),('H3',3)):
            available[tf] = max(0,int((t-opened.ceil(f'{hours}h'))//timedelta(hours=hours)))
        short = names and all(available[requirements[n]['timeframe']] < requirements[n]['lookback'] for n in names)
        reasons.append('SEGMENT_WARMUP' if short else 'SPECIFIC_FEATURE_INCOMPLETE')
        missing.append(json.dumps(names))
    joined['state_diagnostic_reason'],joined['missing_features'] = reasons,missing
    if joined.state_match_status.eq('MATCHED').sum()+joined.state_match_status.eq('NO_COMPLETE_VECTOR_AT_BOUNDARY').sum() != len(trades):
        raise ValueError('state reconciliation failed')
    return joined


def primary_comparisons(joined):
    groups = {s:joined.loc[joined.state_at_entry.eq(s)] for s in STATES}
    days = list(joined.groupby(joined.entry_time.dt.floor('D'),sort=True))
    # Paired day resampling retains intraday dependence and all state populations.
    sufficient = np.zeros((len(days),3,3))
    for i,(_,group) in enumerate(days):
        for j,s in enumerate(STATES):
            g = group.loc[group.state_at_entry.eq(s)]
            sufficient[i,j] = (len(g),g.outcome.eq('WIN').sum(),g.profit.sum())
    rng = np.random.default_rng(42)
    draws = sufficient[rng.integers(len(days),size=(2000,len(days)))].sum(axis=1)
    results = {}
    for other in ('S0','S2'):
        a,b = groups['S1'],groups[other]
        stats = _effect(a.profit.tolist(),b.profit.tolist())
        ai,bi = STATES.index('S1'),STATES.index(other)
        valid = (draws[:,ai,0] > 0) & (draws[:,bi,0] > 0)
        intervals = {}
        for key,col in (('expectancy_difference_eur',2),('win_rate_difference',1)):
            values = draws[valid,ai,col]/draws[valid,ai,0]-draws[valid,bi,col]/draws[valid,bi,0]
            enough = sum(sufficient[:,ai,0]>0) >= 2 and sum(sufficient[:,bi,0]>0) >= 2
            intervals[key] = np.quantile(values,[.025,.975]).tolist() if enough and len(values) >= 2 else None
        results['S1_vs_'+other] = {**stats,
            'win_rate_difference':float(a.outcome.eq('WIN').mean()-b.outcome.eq('WIN').mean()) if len(a) and len(b) else None,
            'S1_expectancy_greater':stats['difference_in_means'] > 0 if stats['difference_in_means'] is not None else None,
            'cluster_bootstrap_95_intervals':intervals,'defined_replicates':int(valid.sum()),
            'entry_day_blocks':{'S1':int(sum(sufficient[:,ai,0]>0)),other:int(sum(sufficient[:,bi,0]>0))}}
    return results


def summarize(joined):
    matched = joined.loc[joined.state_match_status.eq('MATCHED')]
    unmatched = joined.loc[joined.state_match_status.ne('MATCHED')]
    comparisons = primary_comparisons(joined)
    flags = [r['S1_expectancy_greater'] for r in comparisons.values()]
    monthly = []
    for month in range(1,13):
        f = joined.loc[joined.entry_time.dt.month.eq(month)]
        monthly.append({'month':f'2025-{month:02d}','overall':h5.performance(f),
            'matched':int(f.state_match_status.eq('MATCHED').sum()),
            'unmatched':int(f.state_match_status.ne('MATCHED').sum()),
            'states':{s:h5.performance(f.loc[f.state_at_entry.eq(s)]) for s in STATES}})
    missing = Counter(name for text in unmatched.missing_features.dropna() for name in json.loads(text))
    return {'overall':h5.performance(joined),'reconciliation':{'total':len(joined),'matched':len(matched),'unmatched':len(unmatched)},
        'by_state':{s:h5.performance(matched.loc[matched.state_at_entry.eq(s)]) for s in STATES},
        'unmatched_results':h5.performance(unmatched),'primary_comparisons':comparisons,
        'directional_ordering_reproduced':None if any(f is None for f in flags) else all(flags),
        'materiality_verdict':'No frozen numeric materiality threshold; assess reported differences and uncertainty descriptively.',
        'missing_state_reasons':dict(Counter(unmatched.state_diagnostic_reason)),
        'missing_feature_frequency':dict(missing),'monthly':monthly}


def run():
    targets = (OUTPUT/'h7_holdout_2025_trades.csv',OUTPUT/'h7_holdout_2025_summary.json')
    attempt_path = OUTPUT/'h7_holdout_2025.attempt.json'
    if attempt_path.exists() or any(p.exists() for p in targets):
        raise FileExistsError('H7 attempt/output already exists; one-shot execution cannot be repeated')
    protected = {p for p in OUTPUT.rglob('*') if p.is_file()} | {REPORT,DATA,h5.EA}
    before = {str(p):sha256_file(p) for p in sorted(protected)}
    if before[str(h5.EA)] != h5.EA_SHA:
        raise ValueError('canonical EA SHA mismatch')
    with no_fitting():
        contract,model,contract_file_sha = load_frozen_inputs()
        profile,policy,continuity = frozen_continuity(contract)
    feature_metadata = Path(contract['feature_artifact']['path']).with_name('metadata.json')
    meta = json.loads(feature_metadata.read_text())
    if meta['definitions'] != [asdict(d) for d in definitions('M15')]:
        raise ValueError('feature definitions differ from frozen artifact')
    before[str(feature_metadata)] = sha256_file(feature_metadata)
    frozen_code = [*Path('sandbox/market_state/universal').glob('*.py'),Path(h5.__file__),
                   Path('sandbox/market_data/gaps.py'),Path('sandbox/market_data/sanitation.py')]
    before.update({str(p):sha256_file(p) for p in frozen_code})
    attempt = {'status':'RUNNING','protocol':PROTOCOL,'input_sha256':{str(p):before[str(p)] for p in (REPORT,DATA)},
               'implementation_sha256':sha256_file(Path(__file__))}
    with attempt_path.open('x',encoding='utf-8') as stream:
        stream.write(json.dumps(attempt,indent=2,sort_keys=True)+'\n')
    try:
        with no_fitting() as guard:
            trades,report = read_report(REPORT)
            bars,gaps,data = read_bars(DATA,profile,policy)
            if trades.entry_time.min() < bars.timestamp_utc.min() or trades.exit_time.max() > bars.timestamp_utc.max()+STEP:
                raise ValueError('deal timestamps outside source coverage')
            features,states = state_rows(bars,model)
            joined = attach(trades,states,features,bars,meta['definitions'])
            summary = summarize(joined)
        if guard['fit_calls']:
            raise RuntimeError('fit attempted')
        if any(sha256_file(Path(p)) != digest for p,digest in before.items()):
            raise RuntimeError('protected input changed')
        summary.update(experiment=PROTOCOL['version'],protocol=PROTOCOL,data=data,MT5_report=report,
            continuity=continuity,feature_rows=len(features),complete_state_rows=len(states),
            provenance={'protected_sha256_before_and_after':before,'fit_calls':guard['fit_calls'],
                'model_sha256':sha256_file(h5.DEFAULT_MODEL),'contract_canonical_sha256':CONTRACT_SHA256,
                'contract_file_sha256':contract_file_sha,'ea_sha256':sha256_file(h5.EA),
                'report_sha256':sha256_file(REPORT),'csv_sha256':sha256_file(DATA)},
            limitations=['Materially better has no prespecified numerical threshold; no post-hoc cutoff is introduced.',
                'Pointwise day-cluster bootstrap intervals do not establish independence across days or remove multiple-comparison concerns.',
                'State-ineligible trades remain separate; no state or history is imputed.',
                'Tester history-quality label is recorded, not independent certification of execution fidelity. H8 remains separate.',
                'Holdout file begins in 2025; no additional pre-period history was introduced or synthesized.'])
        csv = joined.to_csv(index=False,lineterminator='\n')
        summary['provenance']['trades_csv_sha256'] = hashlib.sha256(csv.encode()).hexdigest()
        for target,payload in zip(targets,(csv,json.dumps(summary,indent=2,sort_keys=True,allow_nan=False)+'\n')):
            with target.open('x',encoding='utf-8',newline='') as stream:stream.write(payload)
        attempt.update(status='COMPLETED',fit_calls=guard['fit_calls'],outputs={str(p):sha256_file(p) for p in targets})
    except Exception as error:
        attempt.update(status='FAILED',error_type=type(error).__name__,error=str(error))
        raise
    finally:
        attempt_path.write_text(json.dumps(attempt,indent=2,sort_keys=True)+'\n',encoding='utf-8')
    return summary


if __name__ == '__main__':
    report = run()
    print(json.dumps({k:report[k] for k in ('overall','reconciliation','by_state','primary_comparisons',
        'directional_ordering_reproduced','missing_state_reasons')},indent=2))
