"""January 2016 ledger-to-frozen-state pipeline proof; never executes a strategy."""
from __future__ import annotations

import argparse
from decimal import Decimal
from datetime import timedelta
import hashlib
import json
from pathlib import Path

import numpy as np
import pandas as pd
import pyarrow.parquet as pq
from openpyxl import load_workbook

from sandbox.provenance import sha256_file
from sandbox.market_data.sanitation import frame_hash
from sandbox.market_state.universal.market_state_v1 import MARKET_STATE_FEATURE_NAMES_V1 as FEATURES
from .future_behavior import _guard
from .validation import load_frozen_inputs, no_fitting, DEFAULT_MODEL, DEFAULT_CONTRACT
from .validation_contract import CONTRACT_SHA256

START = pd.Timestamp('2016-01-01', tz='UTC')
END = pd.Timestamp('2016-02-01', tz='UTC')
STEP = timedelta(minutes=15)
STATES = ('S0', 'S1', 'S2')
EA = Path('reference_strategies/F1_NY3H_Green_MT5_Version0.mq5')
EA_SHA = '80736cd520f5a28836bdf271e5a01036f1e2c3fb4f208e42d612f32ada08ce05'
METRICS = ('posterior_entropy', 'max_posterior', 'posterior_margin', 'state_age', 'bars_since_transition')
OUTPUT = Path('results/market_state')


def parse_report_rows(rows):
    """Pair actual deals only, enforcing the canonical one-position ledger shape."""
    rows = [tuple(r) for r in rows]
    def setting(name):
        matches = [r[3] for r in rows if r[0] == name]
        if len(matches) != 1:
            raise ValueError(f'missing/ambiguous report field {name}')
        return matches[0]
    for key, value in {'Expert:': 'F1_NY3H_Green_MT5_Version0', 'Symbol:': 'EURUSD_DUKASCOPY',
                       'Period:': 'M15 (2016.01.01 - 2016.02.01)', 'Currency:': 'EUR'}.items():
        if setting(key) != value:
            raise ValueError(f'noncanonical report {key}')
    required = {'TakeProfitPips=5', 'StopLossPips=40', 'PipSize=0.0001', 'FixedLots=0.01',
                'MagicNumber=26082930', 'EnableTrading=true', 'UseManualServerUTCOffset=true',
                'ManualServerUTCOffsetHours=0'}
    settings = [r[3] for r in rows if len(r) > 3 and isinstance(r[3], str) and '=' in r[3]]
    if set(settings) != required or len(settings) != len(required):
        raise ValueError('canonical EA inputs mismatch')
    headers = [i for i, r in enumerate(rows) if r[:5] == ('Time', 'Deal', 'Symbol', 'Type', 'Direction')]
    if len(headers) != 1:
        raise ValueError('unique Deals table required')
    trades, active, seen = [], None, set()
    previous = START
    initial = Decimal(str(setting('Initial Deposit:')))
    balance = initial
    def money(x):
        value = Decimal(str(x))
        if not value.is_finite():
            raise ValueError('nonfinite money')
        return value
    for r in rows[headers[0]+1:]:
        if not any(x is not None for x in r):
            continue
        timestamp = pd.to_datetime(r[0], format='%Y.%m.%d %H:%M:%S', utc=True)
        if not START <= timestamp < END or timestamp < previous or r[1] in seen:
            raise ValueError('invalid/duplicate/out-of-window deal')
        previous = timestamp
        seen.add(r[1])
        if r[3] == 'balance':
            if len(seen) != 1 or money(r[10]) != initial or money(r[11]) != initial:
                raise ValueError('unexpected balance operation')
            continue
        if r[2] != 'EURUSD_DUKASCOPY' or money(r[5]) != Decimal('.01'):
            raise ValueError('symbol/volume mismatch')
        price = float(r[6])
        if not np.isfinite(price) or price <= 0:
            raise ValueError('invalid fill price')
        cash = sum((money(r[k]) for k in (8, 9, 10)), Decimal(0))
        balance += cash
        if balance != money(r[11]):
            raise ValueError('deal balance reconciliation failed')
        if (r[3], r[4]) == ('buy', 'in') and active is None:
            active = {'entry_time': timestamp, 'entry_price': price, 'entry_deal': r[1],
                      'entry_order': r[7], 'entry_cash': cash}
        elif (r[3], r[4]) == ('sell', 'out') and active is not None:
            profit = active.pop('entry_cash') + cash
            if profit == 0:
                raise ValueError('unexpected zero-profit January trade')
            trades.append({**active, 'exit_time': timestamp, 'exit_price': price,
                           'exit_deal': r[1], 'profit': float(profit),
                           'deal_profit': float(money(r[10])), 'swap': float(money(r[9])),
                           'commission': float(money(r[8])), 'outcome': 'WIN' if profit > 0 else 'LOSS'})
            active = None
        else:
            raise ValueError('unsupported partial/overlapping/noncanonical deal sequence')
    if active is not None or not trades:
        raise ValueError('incomplete position ledger')
    frame = pd.DataFrame(trades)
    wins = int(frame.outcome.eq('WIN').sum())
    net = balance - initial
    if (len(frame), wins, len(frame)-wins, net) != (49, 43, 6, Decimal('1.22')):
        raise ValueError('expected January reconciliation failed')
    if money(setting('Total Net Profit:')) != net or setting('Total Trades:') != 49 or setting('Total Deals:') != 98:
        raise ValueError('report summary reconciliation failed')
    return frame, {'passed': True, 'trades': 49, 'wins': 43, 'losses': 6, 'win_rate_percent': 100*43/49,
                   'net_profit_eur': float(net), 'history_quality': setting('History Quality:'),
                   'profit_definition': 'net EUR cash: deal profit + swap + commission on entry and exit',
                   'time_source': 'actual Deals in/out timestamps, UTC; never pending-order creation'}


def read_report(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError('expected single MT5 report sheet')
        return parse_report_rows(workbook.worksheets[0].iter_rows(values_only=True))
    finally:
        workbook.close()


def read_features(path, contract):
    path = Path(path)
    if sha256_file(path) != contract['feature_artifact']['sha256']:
        raise ValueError('frozen feature artifact checksum mismatch')
    schema = pq.read_schema(path)
    _guard(schema.names)
    if tuple(c for c in schema.names if c in FEATURES) != FEATURES:
        raise ValueError('feature ordering mismatch')
    columns = ['timestamp', 'continuity_segment_id', 'provider', 'symbol', 'decision_timeframe', *FEATURES]
    return pd.read_parquet(path, columns=columns, filters=[('timestamp', '>=', START), ('timestamp', '<', END)])


def assign_states(features, model):
    """Only complete frozen vectors. Ages reset at gaps, segments and missing vectors."""
    _guard(features.columns)
    if tuple(c for c in features if c in FEATURES) != FEATURES:
        raise ValueError('feature ordering mismatch')
    f = features.copy().sort_values('timestamp', kind='stable').reset_index(drop=True)
    t = f.timestamp
    if (not isinstance(t.dtype, pd.DatetimeTZDtype) or t.isna().any() or t.duplicated().any()
            or not t.between(START, END, inclusive='left').all() or not t.eq(t.dt.floor('15min')).all()):
        raise ValueError('unique aware January M15 close timestamps required')
    if f.continuity_segment_id.isna().any():
        raise ValueError('missing continuity identity')
    for name, value in [('provider', 'DUKASCOPY'), ('symbol', 'EURUSD'), ('decision_timeframe', 'M15')]:
        if not f[name].eq(value).all():
            raise ValueError('feature source mismatch')
    f = f.dropna(subset=list(FEATURES)).reset_index(drop=True)
    if f.empty or not np.isfinite(f.loc[:, FEATURES].to_numpy(float)).all():
        raise ValueError('no finite complete feature vectors')
    with no_fitting():
        x = f.loc[:, FEATURES].to_numpy(float)
        labels = model.predict(x, feature_names=FEATURES)
        p = model.predict_proba(x, feature_names=FEATURES)
    if (p.shape != (len(f), 3) or not np.isfinite(p).all() or (p < 0).any() or (p > 1).any()
            or not np.allclose(p.sum(axis=1), 1, atol=1e-12, rtol=0)
            or not np.array_equal(labels, p.argmax(axis=1))):
        raise ValueError('invalid frozen state outputs')
    connected = (f.timestamp.diff().eq(STEP) & f.continuity_segment_id.eq(f.continuity_segment_id.shift())).to_numpy()
    age = np.ones(len(f), dtype=int)
    since = np.full(len(f), np.nan)
    for i in range(1, len(f)):
        if connected[i]:
            if labels[i] != labels[i-1]:
                since[i] = 0
            else:
                age[i] = age[i-1] + 1
                since[i] = since[i-1] + 1
    logp = np.zeros_like(p)
    np.log(p, out=logp, where=p > 0)
    ordered = np.sort(p, axis=1)
    return pd.DataFrame({'state_timestamp': f.timestamp, 'state_at_entry': [STATES[int(s)] for s in labels],
                         'p_S0': p[:, 0], 'p_S1': p[:, 1], 'p_S2': p[:, 2],
                         'max_posterior': p.max(axis=1), 'posterior_entropy': -(p*logp).sum(axis=1),
                         'posterior_margin': ordered[:, 2]-ordered[:, 1], 'state_age': age,
                         'bars_since_transition': since, 'continuity_segment_id': f.continuity_segment_id})


def attach(trades, states):
    """Exact preceding completed M15 boundary; never carry a stale state across gaps."""
    result = trades.copy()
    result['required_state_timestamp'] = result.entry_time.dt.floor('15min')
    result = result.merge(states, how='left', left_on='required_state_timestamp', right_on='state_timestamp',
                          validate='many_to_one', sort=False)
    result['state_match_status'] = np.where(result.state_at_entry.notna(), 'MATCHED', 'NO_COMPLETE_VECTOR_AT_BOUNDARY')
    if (result.state_timestamp > result.entry_time).any():
        raise ValueError('future state at entry')
    return result


def performance(frame):
    n = len(frame)
    wins = int(frame.outcome.eq('WIN').sum())
    return {'N': n, 'wins': wins, 'losses': int(frame.outcome.eq('LOSS').sum()),
            'WR': wins/n if n else None, 'expectancy_eur': float(frame.profit.mean()) if n else None,
            'net_profit_eur': round(float(frame.profit.sum()), 2)}


def comparisons(frame):
    result = {}
    for outcome in ('WIN', 'LOSS'):
        group = frame.loc[frame.outcome.eq(outcome)]
        result[outcome] = {'N': len(group), 'state_distribution': {
            s: int(group.state_at_entry.fillna('UNMATCHED').eq(s).sum()) for s in (*STATES, 'UNMATCHED')},
            'metrics': {}}
        for name in METRICS:
            values = group[name].dropna()
            result[outcome]['metrics'][name] = {'n': len(values), 'missing': len(group)-len(values),
                'mean': float(values.mean()) if len(values) else None,
                'median': float(values.median()) if len(values) else None}
    return result


def bootstrap(frame, replicates=2000):
    """Fixed UTC-entry-day clusters, whole trades retained. No outcome-based block tuning."""
    groups = [g for _, g in frame.groupby(frame.entry_time.dt.floor('D'), sort=True)]
    rng = np.random.default_rng(42)
    scopes = ('ALL', *STATES, 'UNMATCHED')
    samples = {s: {'WR': [], 'expectancy_eur': []} for s in scopes}
    differences = {m: [] for m in METRICS}
    for _ in range(replicates if len(groups) >= 2 else 0):
        sample = pd.concat([groups[i] for i in rng.integers(len(groups), size=len(groups))], ignore_index=True)
        for s in scopes:
            subset = sample if s == 'ALL' else sample.loc[sample.state_at_entry.fillna('UNMATCHED').eq(s)]
            if len(subset):
                samples[s]['WR'].append(float(subset.outcome.eq('WIN').mean()))
                samples[s]['expectancy_eur'].append(float(subset.profit.mean()))
        for name in METRICS:
            win = sample.loc[sample.outcome.eq('WIN'), name].dropna()
            loss = sample.loc[sample.outcome.eq('LOSS'), name].dropna()
            if len(win) and len(loss):
                differences[name].append(float(win.mean()-loss.mean()))
    def interval(values, blocks):
        return {'contributing_blocks': int(blocks), 'defined_replicates': len(values),
                'percentile_95_ci': [float(x) for x in np.quantile(values, [.025, .975])]
                if blocks >= 2 and len(values) >= 2 else None}
    return {'method': 'whole UTC entry-day cluster bootstrap; 95% pointwise percentile intervals',
            'seed': 42, 'replicates': replicates, 'entry_day_blocks': len(groups),
            'performance': {s: {m: interval(v, sum(len(g if s == 'ALL' else g.loc[g.state_at_entry.fillna('UNMATCHED').eq(s)]) > 0 for g in groups))
                               for m, v in values.items()} for s, values in samples.items()},
            'winner_minus_loser_mean': {m: interval(v, min(sum(g.loc[g.outcome.eq(o), m].notna().any() for g in groups)
                                                          for o in ('WIN', 'LOSS'))) for m, v in differences.items()},
            'limitations': 'Very few loss days. Approximate independence across days is unverified; cross-day dependence and multiplicity remain. No inferential H5 conclusion.'}


def run(report=Path('sandbox/test.xlsx'), output=OUTPUT):
    output = Path(output)
    targets = [output/'h5_january_2016_trades.csv', output/'h5_january_2016_summary.json']
    if any(p.exists() for p in targets):
        raise FileExistsError('H5 output already exists; refusing overwrite')
    protected = set(OUTPUT.glob('*.json')) | {DEFAULT_MODEL, DEFAULT_CONTRACT, EA, Path(report)}
    before = {str(p): sha256_file(p) for p in sorted(protected)}
    if before[str(EA)] != EA_SHA:
        raise ValueError('canonical EA checksum mismatch')
    with no_fitting() as guard:
        contract, model, contract_file_sha = load_frozen_inputs()
        features = read_features(contract['feature_artifact']['path'], contract)
        states = assign_states(features, model)
        trades, reconciliation = read_report(report)
        joined = attach(trades, states)
    if guard['fit_calls']:
        raise RuntimeError('fit attempted')
    missing = int(joined.state_at_entry.isna().sum())
    summary = {'experiment': 'H5_JANUARY_2016_PIPELINE_PROOF_V1',
        'pipeline_status': 'PASSED' if not missing else 'FAILED_INCOMPLETE_STATE_COVERAGE',
        'reconciliation': reconciliation, 'overall': performance(joined),
        'by_state': {s: performance(joined.loc[joined.state_at_entry.eq(s)]) for s in STATES},
        'unmatched': performance(joined.loc[joined.state_at_entry.isna()]),
        'missing_state_matches': missing, 'winner_loser': comparisons(joined), 'bootstrap': bootstrap(joined),
        'provenance': {'report_path': str(report), 'report_sha256': before[str(Path(report))],
            'frozen_model_sha256': before[str(DEFAULT_MODEL)], 'contract_checksum': CONTRACT_SHA256,
            'contract_file_sha256': contract_file_sha, 'canonical_ea_sha256': before[str(EA)],
            'feature_artifact': contract['feature_artifact'], 'feature_names': list(FEATURES),
            'january_feature_rows': len(features), 'eligible_state_observations': len(states),
            'january_features_sha256': frame_hash(features), 'assigned_states_sha256': frame_hash(states),
            'fit_calls': guard['fit_calls'], 'protected_before': before},
        'definitions': {'state_timestamp': 'Feature decision/close time; exact floor(entry UTC, 15 minutes). No forward or stale gap match.',
            'state_age': 'One-based observed consecutive-state M15 bars; resets at gap, missing vector or segment change. Initial age left-censored.',
            'bars_since_transition': 'Zero on observed transition; increments on continuation; unknown until first transition after a gap.',
            'entropy': 'Natural logarithm; zero probability contributes zero.',
            'profit': 'Net EUR cash including swap and commission; no reconstructed fills.'},
        'limitations': ['Pipeline proof only: 49 trades / 6 losses, no final H5 conclusion, thresholds or filters.',
            'Frozen model was fitted on 2016-2022: January assignment is retrospective/in-sample, not a model available live in January 2016. Feature information is entry-causal.',
            'MT5 report states 93% real ticks; full native-tick execution fidelity is not established by this report.',
            'Trade feed and frozen historical feature feed are both Dukascopy but exact quote-stream equivalence is not established.',
            'Missing warm-up/gap states are retained as unmatched, never imputed or assigned from future observations.']}
    after = {p: sha256_file(Path(p)) for p in before}
    if before != after:
        raise RuntimeError('protected input changed')
    summary['provenance']['protected_after'] = after
    csv = joined.to_csv(index=False, lineterminator='\n').encode('utf-8')
    summary['provenance']['trades_csv_sha256'] = hashlib.sha256(csv).hexdigest()
    encoded = (json.dumps(summary, sort_keys=True, indent=2, allow_nan=False)+'\n').encode()
    output.mkdir(parents=True, exist_ok=True)
    for path, data in zip(targets, (csv, encoded)):
        with path.open('xb') as stream:
            stream.write(data)
    return summary


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--report', type=Path, default=Path('sandbox/test.xlsx'))
    args = parser.parse_args()
    result = run(args.report)
    print(json.dumps({k: result[k] for k in ('pipeline_status', 'overall', 'by_state', 'missing_state_matches')}, indent=2))
