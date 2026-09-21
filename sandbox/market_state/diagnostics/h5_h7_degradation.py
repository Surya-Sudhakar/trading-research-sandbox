"""Post-H7 descriptive accounting/drift diagnostics; 2025 is exposed data.

No experiment is rerun or updated. Saved ledgers control membership and outcomes.
Observation diagnostics replay existing frozen inference, requiring full ledger parity.
"""
from __future__ import annotations

from dataclasses import asdict
import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from sandbox.market_state.discovery import h5_discovery_2016_2022 as h5
from sandbox.market_state.discovery.dukascopy_h5_coverage import frozen_continuity
from sandbox.market_state.discovery.future_behavior import _effect
from sandbox.market_state.discovery.validation import load_frozen_inputs, no_fitting
from sandbox.market_state.universal.schema import definitions
from sandbox.market_state.validation import h7_holdout_2025 as h7
from sandbox.provenance import sha256_file

ROOT = Path('results/market_state')
STATES = ('S0', 'S1', 'S2')
VARIABLES = ('state_age', 'bars_since_transition', 'max_posterior',
             'posterior_margin', 'posterior_entropy')
EXPECTED = {'H5': (4162, 2384, (114, 794, 1476)), 'H7': (347, 234, (15, 63, 156))}
PREFIX = {'H5': 'h5_discovery_2016_2022', 'H7': 'h7_holdout_2025'}
OUTPUTS = ('h5_h7_degradation_summary.json', 'h5_h7_degradation_by_state.csv',
           'h5_h7_s1_distribution_drift.csv')


def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def check_ledger(frame, era):
    """Fail closed on population, UTC boundary, state, or cash inconsistencies."""
    f = frame.copy(deep=True)
    for col in ('entry_time', 'exit_time', 'required_state_timestamp', 'state_timestamp'):
        f[col] = pd.to_datetime(f[col], utc=True)
    start, end = ('2016-01-01', '2023-01-01') if era == 'H5' else ('2025-01-01', '2026-01-01')
    # H5's frozen membership is by entry date; one unmatched December trade exits
    # in January 2023. Preserve it rather than silently changing the denominator.
    for col in (('entry_time',) if era == 'H5' else ('entry_time', 'exit_time')):
        if not f[col].between(pd.Timestamp(start, tz='UTC'), pd.Timestamp(end, tz='UTC'), inclusive='left').all():
            raise ValueError('trade outside diagnostic period')
    if not f.required_state_timestamp.eq(f.entry_time.dt.floor('15min')).all():
        raise ValueError('changed causal boundary')
    matched = f.state_match_status.eq('MATCHED')
    total, count, state_counts = EXPECTED[era]
    if (len(f) != total or int(matched.sum()) != count
            or tuple(int((matched & f.state_at_entry.eq(s)).sum()) for s in STATES) != state_counts
            or not f.state_match_status.isin(['MATCHED', 'NO_COMPLETE_VECTOR_AT_BOUNDARY']).all()):
        raise ValueError('frozen ledger reconciliation failed')
    if (f.exit_time.isna().any() or (f.exit_time<f.entry_time).any()
            or f.entry_deal.duplicated().any() or f.exit_deal.duplicated().any()
            or not f.loc[matched, 'state_timestamp'].eq(f.loc[matched, 'required_state_timestamp']).all()
            or f.loc[~matched, 'state_at_entry'].notna().any()
            or not np.isfinite(f.profit).all()
            or not f.outcome.eq(np.where(f.profit > 0, 'WIN', 'LOSS')).all()
            or f.profit.eq(0).any()):
        raise ValueError('invalid saved ledger')
    return f


def cash_from_deals(frame, deals):
    """Both sides of actual deals; never infer entry costs from exit-only columns."""
    result = frame.copy(deep=True)
    totals = []
    for row in frame.itertuples():
        pair = [deals[int(row.entry_deal)], deals[int(row.exit_deal)]]
        values = np.sum(pair, axis=0)  # gross profit, commission, swap
        if not np.isclose(values.sum(), row.profit, rtol=0, atol=1e-8):
            raise ValueError('report cash does not reconcile with frozen ledger')
        totals.append(values)
    result[['gross_profit', 'commission_total', 'swap_total']] = np.asarray(totals)
    return result


def report_cash(path, frame):
    """Read headers and only the selected frozen ledger deal IDs; no extra periods."""
    wanted = set(frame.entry_deal.astype(int)) | set(frame.exit_deal.astype(int))
    details = {}
    for trade in frame.itertuples():
        details[int(trade.entry_deal)] = ('buy','in',trade.entry_time,trade.entry_price)
        details[int(trade.exit_deal)] = ('sell','out',trade.exit_time,trade.exit_price)
    deals, settings = {}, []
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        active, header = False, True
        for row in workbook.active.iter_rows(values_only=True):
            if row[0] == 'Deals':
                active = True
                continue
            if not active:
                if row[0] == 'Results':
                    header = False
                if header or row[0] in ('History Quality:', 'Bars:'):
                    settings.append([str(v) for v in row if v is not None])
                continue
            try:
                deal_id = int(row[1])
            except (TypeError, ValueError):
                continue
            if deal_id not in wanted:
                continue
            if deal_id in deals:
                raise ValueError('duplicate report deal')
            side,direction,timestamp,price = details[deal_id]
            if ((row[3],row[4]) != (side,direction) or pd.Timestamp(row[0],tz='UTC') != timestamp
                    or not np.isclose(float(row[6]),price,atol=1e-12,rtol=0)
                    or not np.isclose(float(row[5]),.01,atol=1e-12,rtol=0)):
                raise ValueError('report deal identity mismatch')
            deals[deal_id] = tuple(float(row[i] or 0) for i in (10, 8, 9))
    finally:
        workbook.close()
    if set(deals) != wanted:
        raise ValueError('missing canonical report deals')
    flat = {v for row in settings for v in row}
    if not set(h7.INPUTS).issubset(flat) or 'F1_NY3H_Green_MT5_Version0' not in flat or 'EUR' not in flat:
        raise ValueError('canonical report settings mismatch')
    return cash_from_deals(frame, deals), settings


def distribution(values):
    a = np.asarray(pd.Series(values).dropna(), dtype=float)
    if not np.isfinite(a).all():
        raise ValueError('nonfinite diagnostic input')
    result = {'N': len(a), 'missing': len(values)-len(a)}
    for name in ('mean', 'median', 'std', 'q25', 'q75', 'IQR'):
        result[name] = None
    if len(a):
        q25, med, q75 = np.quantile(a, [.25, .5, .75])
        result.update(mean=float(a.mean()), median=float(med), q25=float(q25), q75=float(q75),
                      IQR=float(q75-q25), std=float(a.std(ddof=1)) if len(a)>1 else None)
    return result


def drift(before, after):
    """H7 minus H5; empirical two-sample KS without IID p-values or drift cutoffs."""
    a, b = np.sort(pd.Series(before).dropna().to_numpy(float)), np.sort(pd.Series(after).dropna().to_numpy(float))
    left, right = distribution(before), distribution(after)
    result = {'H5': left, 'H7': right, 'mean_difference': None, 'median_difference': None,
              'IQR_difference': None, 'standardized_mean_difference': None, 'KS': None}
    if len(a) and len(b):
        grid = np.unique(np.concatenate([a, b]))
        result.update(mean_difference=right['mean']-left['mean'], median_difference=right['median']-left['median'],
                      IQR_difference=right['IQR']-left['IQR'],
                      KS=float(np.max(np.abs(np.searchsorted(a, grid, side='right')/len(a)
                                             - np.searchsorted(b, grid, side='right')/len(b)))))
        if len(a)>1 and len(b)>1:
            # Equal population weighting avoids letting the seven-year sample set the scale alone.
            sd = np.sqrt((a.var(ddof=1)+b.var(ddof=1))/2)
            result['standardized_mean_difference'] = float((b.mean()-a.mean())/sd) if sd else None
    return result


def payoffs(frame, value='profit'):
    wins = frame.loc[frame.outcome.eq('WIN'), value]
    losses = frame.loc[frame.outcome.eq('LOSS'), value]
    return {**h5.performance(frame), 'mean_win': float(wins.mean()) if len(wins) else None,
            'median_win': float(wins.median()) if len(wins) else None,
            'mean_loss': float(losses.mean()) if len(losses) else None,
            'median_loss': float(losses.median()) if len(losses) else None,
            'payoff_ratio': float(wins.mean()/abs(losses.mean())) if len(wins) and len(losses) and losses.mean() else None}


def decompose(before, after, value='profit'):
    """Exact symmetric decomposition; no arbitrary first/second substitution order."""
    a, b = payoffs(before, value), payoffs(after, value)
    if any(d[k] is None for d in (a, b) for k in ('mean_win', 'mean_loss')):
        raise ValueError('both outcomes required for decomposition')
    p = (a['WR']+b['WR'])/2
    contributions = {
        'win_probability': (b['WR']-a['WR'])*((a['mean_win']+b['mean_win'])/2-(a['mean_loss']+b['mean_loss'])/2),
        'average_winner': p*(b['mean_win']-a['mean_win']),
        'average_loser': (1-p)*(b['mean_loss']-a['mean_loss'])}
    delta = float(after[value].mean()-before[value].mean())
    if not np.isclose(sum(contributions.values()), delta, atol=1e-12, rtol=0):
        raise ValueError('expectancy decomposition failed')
    return {'change_eur_per_trade': delta, 'contributions': contributions}


def verify_state_parity(ledger, states):
    replay = h5.attach(ledger[['entry_time']].copy(), states)
    for name in ('state_match_status', 'state_at_entry'):
        if not replay[name].fillna('').equals(ledger[name].reset_index(drop=True).fillna('')):
            raise ValueError('frozen state membership mismatch')
    for name in ('p_S0', 'p_S1', 'p_S2', *VARIABLES):
        if not np.allclose(replay[name], ledger[name], atol=1e-12, rtol=0, equal_nan=True):
            raise ValueError('frozen state numeric mismatch: '+name)


def observation_summary(states):
    f = states.sort_values('state_timestamp', kind='stable').reset_index(drop=True)
    connected = (f.state_timestamp.diff().eq(h5.STEP)
                 & f.continuity_segment_id.eq(f.continuity_segment_id.shift())).to_numpy()
    labels = f.state_at_entry.map({s:i for i,s in enumerate(STATES)}).to_numpy()
    counts = np.zeros((3,3), dtype=int)
    np.add.at(counts, (labels[:-1][connected[1:]], labels[1:][connected[1:]]), 1)
    starts = np.flatnonzero(~connected | np.r_[True, labels[1:] != labels[:-1]])
    lengths = np.diff(np.r_[starts,len(f)])
    return {'N': len(f), 'eligible_contiguous_blocks': int((~connected).sum()),
            'connected_pairs': int(connected.sum()), 'transition_counts': counts.tolist(),
            'transition_probabilities': [[float(v/row.sum()) for v in row] if row.sum() else [None]*3 for row in counts],
            'state_change_fraction_of_connected_pairs': float((counts.sum()-np.trace(counts))/counts.sum()) if counts.sum() else None,
            'states': {s: {'count':int((labels==i).sum()), 'fraction':float((labels==i).mean()),
                           'run_lengths_bars': distribution(lengths[labels[starts]==i]),
                           'metrics':{v:distribution(f.loc[f.state_at_entry.eq(s),v]) for v in VARIABLES}}
                       for i,s in enumerate(STATES)}}


def entry_features(ledger, features):
    matched = ledger.loc[ledger.state_match_status.eq('MATCHED')].copy()
    result = matched.merge(features[['timestamp', *h5.FEATURES]], how='left',
                           left_on='required_state_timestamp', right_on='timestamp', validate='many_to_one')
    if result[list(h5.FEATURES)].isna().any().any():
        raise ValueError('saved matched trade lacks complete exact-boundary feature row')
    return result


def calendar_groups(frame):
    result = {}
    for name, values, keys in (('hour_utc',frame.entry_time.dt.hour,range(24)),
                               ('weekday',frame.entry_time.dt.weekday,range(7)),
                               ('month',frame.entry_time.dt.month,range(1,13))):
        loss_total = int(frame.outcome.eq('LOSS').sum())
        result[name] = [{name:int(k), **h5.performance(frame.loc[values.eq(k)]),
                         'share_of_S1_trades':int(values.eq(k).sum())/len(frame),
                         'share_of_S1_losses':int((values.eq(k)&frame.outcome.eq('LOSS')).sum())/loss_total if loss_total else None}
                        for k in keys]
    return result


def mean_uncertainty(frame):
    """Fixed UTC-entry-day cluster bootstrap; descriptive, not a decision criterion."""
    groups = frame.groupby(frame.entry_time.dt.floor('D'), sort=True).profit.agg(['count','sum'])
    if len(groups)<2:
        return {'entry_day_blocks':len(groups), 'interval_95':None}
    a = groups.to_numpy(float)
    rng = np.random.default_rng(42)
    draws = a[rng.integers(len(a),size=(2000,len(a)))].sum(axis=1)
    return {'entry_day_blocks':len(a), 'replicates':2000, 'seed':42,
            'interval_95':np.quantile(draws[:,1]/draws[:,0],[.025,.975]).tolist(),
            'limitation':'Within-day dependence retained; cross-day dependence and feed confounding are not captured.'}


def ensure_unchanged(before):
    for path, digest in before.items():
        if sha256_file(Path(path)) != digest:
            raise ValueError('protected artifact changed: '+path)


def write_outputs(summary, by_state, drift_rows, output=ROOT):
    targets = [output/n for n in OUTPUTS]
    if any(p.exists() for p in targets):
        raise FileExistsError('diagnostic outputs already exist; refusing overwrite')
    # Compute all serialization before creating any output.
    texts = [json.dumps(summary, indent=2, sort_keys=True, allow_nan=False)+'\n',
             pd.DataFrame(by_state).to_csv(index=False), pd.DataFrame(drift_rows).to_csv(index=False)]
    for path, text in zip(targets,texts):
        with path.open('x', encoding='utf-8', newline='') as stream:
            stream.write(text)


def run():
    if any((ROOT/n).exists() for n in OUTPUTS):
        raise FileExistsError('diagnostic already exists')
    protected = {p for p in ROOT.rglob('*') if p.is_file()}
    protected |= set(Path('sandbox/market_state').rglob('*.py')) | {h5.EA}
    protected |= {Path('sandbox/market_data/gaps.py'),Path('sandbox/market_data/sanitation.py')}
    before = {str(p):sha256_file(p) for p in sorted(protected)}
    summaries = {e:read_json(ROOT/(p+'_summary.json')) for e,p in PREFIX.items()}
    ledgers = {}
    for era,prefix in PREFIX.items():
        path = ROOT/(prefix+'_trades.csv')
        if sha256_file(path) != summaries[era]['provenance']['trades_csv_sha256']:
            raise ValueError('saved ledger checksum mismatch')
        ledgers[era] = check_ledger(pd.read_csv(path),era)
    report_paths = {'H5':Path(summaries['H5']['provenance']['report_path']), 'H7':h7.REPORT}
    headers = {}
    for era,path in report_paths.items():
        expected = summaries[era]['provenance']['report_sha256']
        if sha256_file(path) != expected:
            raise ValueError('report checksum mismatch')
        before[str(path)] = expected
        ledgers[era],headers[era] = report_cash(path,ledgers[era])
    if sha256_file(h5.EA) != h5.EA_SHA or sha256_file(h7.DATA) != summaries['H7']['provenance']['csv_sha256']:
        raise ValueError('EA/source checksum mismatch')
    before[str(h7.DATA)] = sha256_file(h7.DATA)
    features,states = {},{}
    with no_fitting() as fit_guard:
        contract,model,contract_file_sha = load_frozen_inputs()
        feature_path = Path(contract['feature_artifact']['path'])
        metadata = read_json(feature_path.with_name('metadata.json'))
        if metadata['definitions'] != [asdict(d) for d in definitions('M15')]:
            raise ValueError('frozen feature definitions changed')
        before[str(feature_path)] = contract['feature_artifact']['sha256']
        before[str(feature_path.with_name('metadata.json'))] = sha256_file(feature_path.with_name('metadata.json'))
        features['H5'] = h5.read_features(feature_path,contract)
        states['H5'] = h5.assign_states(features['H5'],model)
        profile,policy,continuity = frozen_continuity(contract)
        bars,_,_ = h7.read_bars(h7.DATA,profile,policy)
        features['H7'],states['H7'] = h7.state_rows(bars,model)
        if (len(states['H5']),len(states['H7'])) != (93662,16797):
            raise ValueError('eligible observation reconciliation failed')
        for era in ledgers:
            verify_state_parity(ledgers[era],states[era])
    print('Frozen inference replay matches both saved ledgers; fit calls:',fit_guard['fit_calls'],flush=True)
    entries = {e:entry_features(ledgers[e],features[e]) for e in ledgers}
    by_state, drift_rows, decomposition, relative = [],[],{},{}
    for s in STATES:
        groups = {e:entries[e].loc[entries[e].state_at_entry.eq(s)] for e in entries}
        decomposition[s] = {'net':decompose(groups['H5'],groups['H7']),
                             'gross_plus_costs':decompose(groups['H5'],groups['H7'],'gross_profit')}
        alternative = decomposition[s]['gross_plus_costs']
        for name in ('commission_total','swap_total'):
            alternative['contributions'][name] = float(groups['H7'][name].mean()-groups['H5'][name].mean())
        alternative['net_change_eur_per_trade'] = float(sum(alternative['contributions'].values()))
        if not np.isclose(alternative['net_change_eur_per_trade'],decomposition[s]['net']['change_eur_per_trade'],atol=1e-12):
            raise ValueError('gross/cost decomposition failed')
        for era,g in groups.items():
            by_state.append({'era':era,'state':s, **payoffs(g),
                             'matched_trade_share':len(g)/len(entries[era]),
                             'gross_mean_eur':float(g.gross_profit.mean()),
                             'commission_mean_eur':float(g.commission_total.mean()),
                             'swap_mean_eur':float(g.swap_total.mean())})
    observations = {e:observation_summary(states[e]) for e in states}
    for scope, frames in (('S1_eligible_observations', {e:states[e].merge(features[e][['timestamp',*h5.FEATURES]],
                            left_on='state_timestamp',right_on='timestamp',validate='one_to_one') for e in states}),
                          ('S1_matched_entries',entries)):
        frames = {e:f.loc[f.state_at_entry.eq('S1')] for e,f in frames.items()}
        for v in (*h5.FEATURES,*VARIABLES):
            d = drift(frames['H5'][v],frames['H7'][v])
            drift_rows.append({'scope':scope,'variable':v,
                **{f'{e}_{k}':value for e in ('H5','H7') for k,value in d[e].items()},
                **{k:value for k,value in d.items() if k not in ('H5','H7')}})
    within, calendar, yearly, uncertainty, fills = {},{},[],{},{}
    for era,f in entries.items():
        s1 = f.loc[f.state_at_entry.eq('S1')]
        calendar[era] = calendar_groups(s1)
        uncertainty[era] = mean_uncertainty(s1)
        within[era] = {v:{'WIN':distribution(s1.loc[s1.outcome.eq('WIN'),v]),
                          'LOSS':distribution(s1.loc[s1.outcome.eq('LOSS'),v])} for v in (*VARIABLES,*h5.FEATURES)}
        fills[era] = {s:{o:{'signed_price_move_pips':distribution((g.exit_price-g.entry_price)/.0001),
                           'holding_minutes':distribution((g.exit_time-g.entry_time).dt.total_seconds()/60)}
                        for o in ('WIN','LOSS')
                        for g in [f.loc[f.state_at_entry.eq(s)&f.outcome.eq(o)]]} for s in STATES}
        for year in (range(2016,2023) if era=='H5' else (2025,)):
            yearly.append({'year':year,**h5.performance(s1.loc[s1.entry_time.dt.year.eq(year)])})
        relative[era] = {s:_effect(s1.profit.tolist(),f.loc[f.state_at_entry.eq(s),'profit'].tolist()) for s in ('S0','S2')}
    missing5 = read_json(ROOT/'h5_missing_state_diagnostic_summary.json')
    coverage = {}
    for era,ledger in ledgers.items():
        reasons = ({k:v['count'] for k,v in missing5['by_diagnostic_reason'].items()} if era=='H5'
                   else summaries[era]['missing_state_reasons'])
        frequency = missing5['missing_feature_frequency'] if era=='H5' else summaries[era]['missing_feature_frequency']
        n = len(ledger)-len(entries[era])
        coverage[era] = {'total':len(ledger),'matched':len(entries[era]),'unmatched':n,
                         'matched_fraction':len(entries[era])/len(ledger),'reasons':reasons,
                         'missing_feature_frequency':frequency,
                         'reason_fraction_of_unmatched':{k:v/n for k,v in reasons.items()},
                         'missing_feature_fraction_of_unmatched':{k:v/n for k,v in frequency.items()}}
    summary = {'experiment':'POST_H7_DEGRADATION_DIAGNOSTIC_V1','2025_status':'EXPOSED_DATA',
        'D1_comparability':{'report_headers':headers,
          'common':'Canonical EA SHA, TP5/SL40, lots0.01, magic26082930, M15, manual UTC offset0; frozen GMM/features and exact completed boundary; no imputation.',
          'differences':['H5 MT5 EURUSD report versus H7 custom EURUSD_DUKASCOPY symbol; state features use Dukascopy in both.',
                         'H5 report says 0% real ticks; H7 says 100% real ticks. Labels do not establish equivalent intrabar execution or quality.',
                         'H5 entry and exit costs reconstructed: its CSV commission/swap columns alone describe exits. Both net-profit definitions include both sides.',
                         'H5 eligibility is entry-date based: one unmatched December 2022 trade exits in January 2023; H7 requires entry and exit in 2025. No matched state comparison is affected.',
                         'Matching current canonical EA source SHA and report Expert/settings does not independently certify the compiled tester binary.',
                         'H5 frozen saved feature artifact versus H7 frozen pipeline on imported 2025 CSV; no added warm-up in this diagnostic.'],
          'continuity':continuity,'state_replay_parity':True,'report_cash_reconciliation':True,
          'observed_fills_not_reconstructed':fills},
        'D2_by_state':by_state,'D2_decomposition':decomposition,
        'decomposition_definition':'Symmetric exact decomposition: dp*(midpoint mean win - midpoint mean loss), midpoint p*dWin, (1-midpoint p)*dLoss. Net and gross-plus-costs are alternative accountings, not additive. Outcome labels always frozen net WIN/LOSS.',
        'D3_observations':observations,'D3_S1_drift':drift_rows,
        'D3_matched_entry_metrics':{e:{s:{v:distribution(f.loc[f.state_at_entry.eq(s),v]) for v in VARIABLES}
                                      for s in STATES} for e,f in entries.items()},
        'drift_definition':'H7 minus H5, SMD uses sqrt((variance_H5+variance_H7)/2); empirical KS without IID p-values or thresholds. Available-case summaries preserve missingness.',
        'D4_S1_calendar':calendar,'D4_S1_WIN_LOSS':within,'D5_S1_yearly':yearly,
        'D6_relative_effects':relative,'D6_S1_expectancy_day_bootstrap':uncertainty,
        'D6_frozen_H7_relative_uncertainty':summaries['H7']['primary_comparisons'],
        'D7_coverage':coverage,
        'limitations':['Accounting decomposition is not identification of why win probability changed.',
                       'Different feeds/MT5 execution and seven years versus one confound market-distribution and strategy-response changes.',
                       'H5 is discovery/model-development evidence, not pristine external confirmation; selection inflation is plausible, not identified.',
                       'Only 63 S1 trades and 13 losses in exposed 2025; H5 has only 9 S1 losses.',
                       'Dependence, missingness selection and changing feed cannot be removed by bootstrap.',
                       'Relative direction survived; absolute economics did not. Stable exploitable effect is not established.',
                       '2023/2024 are not interpolated; 2025 cannot be reused as an untouched holdout.'],
        'provenance':{'protected_sha256_before_and_after':before,'contract_file_sha256':contract_file_sha,
                      'fit_calls':fit_guard['fit_calls'],'threshold_searches':0,'extra_warmup':False,
                      'implementation_sha256':sha256_file(Path(__file__))}}
    ensure_unchanged(before)
    write_outputs(summary,by_state,drift_rows)
    ensure_unchanged(before)
    print(json.dumps({'outputs':{n:sha256_file(ROOT/n) for n in OUTPUTS},'fit_calls':fit_guard['fit_calls']}),flush=True)


if __name__ == '__main__':
    run()
